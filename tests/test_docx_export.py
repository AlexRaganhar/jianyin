from __future__ import annotations

import zipfile

from docx import Document

from jianyin_app.documents.docx_export import export_docx


def test_export_docx_reflows_headings_tables_and_pages(tmp_path) -> None:
    markdown = (
        "# 标题\n\n"
        "这是第一段。\n\n"
        "| 项目 | 内容 |\n"
        "| --- | --- |\n"
        "| 姓名 | 张某 |\n\n"
        "<!-- PDF page 2 -->\n"
        "第二页正文。\n"
    )
    output = tmp_path / "结果.docx"
    export_docx(markdown, output, title="脱敏文本")
    assert output.is_file()
    document = Document(str(output))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "标题" in text
    assert "第一段" in text
    assert "第二页正文" in text
    assert document.tables
    assert document.tables[0].cell(1, 0).text == "姓名"

    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "张某" in xml
    assert "page 2" not in xml


def test_export_docx_does_not_embed_secret_source(tmp_path) -> None:
    markdown = "姓名 [姓名1]，电话 138*****000。\n"
    output = tmp_path / "脱敏.docx"
    export_docx(markdown, output)
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "Alice" not in xml
    assert "13800138000" not in xml

