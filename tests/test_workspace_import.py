from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from jianyin_app.config import WorkspaceSettings
from jianyin_app.errors import UnsupportedDocumentError, WorkspaceError
from jianyin_app.workspace.store import CaseWorkspace


def _workspace(tmp_path: Path) -> tuple[CaseWorkspace, dict]:
    workspace = CaseWorkspace(tmp_path / "workspace")
    return workspace, workspace.create_case("测试案件")


def _blank_pdf(path: Path, *, password: str | None = None) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    if password:
        writer.encrypt(password)
    with path.open("wb") as stream:
        writer.write(stream)


def test_text_import_preserves_source_and_reads_gb18030(tmp_path) -> None:
    workspace, case = _workspace(tmp_path)
    source = tmp_path / "材料.txt"
    source.write_bytes("张三的电话是13800138000。".encode("gb18030"))
    document = workspace.import_document(case["id"], source)
    assert document["extraction_method"] == "text"
    assert document["source_encoding"] == "gb18030"
    assert "张三" in workspace.read_document_text(case["id"], document["id"])
    assert workspace.resolve_case_path(case["id"], document["source_path"]).is_file()


def test_docx_import_reads_paragraphs_tables_and_headers(tmp_path) -> None:
    workspace, case = _workspace(tmp_path)
    source = tmp_path / "笔录.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "页眉：张三"
    document.add_paragraph("被告人张三")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "证据"
    document.save(source)
    record = workspace.import_document(case["id"], source)
    text = workspace.read_document_text(case["id"], record["id"])
    assert record["extraction_method"] == "docx-text"
    assert "被告人张三" in text
    assert "项目 | 证据" in text
    assert "页眉：张三" in text


def test_pdf_text_import_and_page_marker(tmp_path, monkeypatch) -> None:
    workspace, case = _workspace(tmp_path)
    source = tmp_path / "起诉书.pdf"
    _blank_pdf(source)
    monkeypatch.setattr(
        "jianyin_app.workspace.store.extract_pdf_pages",
        lambda _path: ["第一页 Alice 的联系方式与基本情况说明" * 3, "第二页 Bob 的陈述与相关材料说明" * 3],
    )
    record = workspace.import_document(case["id"], source)
    text = workspace.read_document_text(case["id"], record["id"])
    assert record["extraction_method"] == "pdf-text"
    assert record["page_count"] == 2
    assert "<!-- PDF page 1 -->" in text
    assert "第二页 Bob" in text


def test_encrypted_pdf_uses_password_and_fake_ocr(tmp_path, monkeypatch) -> None:
    workspace, case = _workspace(tmp_path)
    source = tmp_path / "加密.pdf"
    _blank_pdf(source, password="secret")
    monkeypatch.setattr("jianyin_app.workspace.store.extract_pdf_pages", lambda _path: [""])

    def fake_ocr(_input_pdf, output_markdown, workdir):
        pages = workdir / "pages"
        pages.mkdir(parents=True, exist_ok=True)
        (pages / "page_0001.txt").write_text("Alice\n", encoding="utf-8")
        output_markdown.write_text("Alice\n", encoding="utf-8")
        (workdir / "progress.json").write_text(
            json.dumps({"total": 1, "done": [1], "failed": []}),
            encoding="utf-8",
        )

    record = workspace.import_document(case["id"], source, password="secret", ocr_runner=fake_ocr)
    assert record["was_encrypted"] is True
    assert record["extraction_method"] == "pdf-to-markdown-ocr"
    assert "Alice" in workspace.read_document_text(case["id"], record["id"])
    assert workspace.resolve_case_path(case["id"], record["original_encrypted_path"]).is_file()


def test_duplicate_import_is_flagged_but_allowed(tmp_path) -> None:
    workspace, case = _workspace(tmp_path)
    source = tmp_path / "材料.txt"
    source.write_text("Alice\n", encoding="utf-8")
    first = workspace.import_document(case["id"], source)
    second = workspace.import_document(case["id"], source)
    assert first["id"] != second["id"]
    assert second["duplicate_of"] == first["id"]


def test_batch_import_continues_after_failure(tmp_path) -> None:
    workspace, case = _workspace(tmp_path)
    good = tmp_path / "good.txt"
    good.write_text("Alice\n", encoding="utf-8")
    bad = tmp_path / "bad.exe"
    bad.write_bytes(b"not a document")
    outcome = workspace.import_documents(case["id"], [good, bad])
    assert len(outcome.imported) == 1
    assert len(outcome.failures) == 1
    assert outcome.failures[0].code == "DOCUMENT_UNSUPPORTED"


def test_path_traversal_and_settings(tmp_path) -> None:
    workspace, case = _workspace(tmp_path)
    with pytest.raises(WorkspaceError):
        workspace.resolve_case_path(case["id"], "../outside.txt")
    settings = WorkspaceSettings(export_docx=False, window_size=2048, overlap_tokens=128)
    workspace.save_settings(settings)
    assert workspace.settings.export_docx is False
    assert workspace.settings.window_size == 2048


def test_unsupported_file_is_rejected(tmp_path) -> None:
    workspace, case = _workspace(tmp_path)
    source = tmp_path / "bad.exe"
    source.write_bytes(b"x")
    with pytest.raises(UnsupportedDocumentError):
        workspace.import_document(case["id"], source)

