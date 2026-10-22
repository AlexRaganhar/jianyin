"""Local document extraction for TXT, Markdown, DOCX, and PDF."""

from __future__ import annotations

import hashlib
import re
import shutil
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader, PdfWriter

from ..errors import DocumentError, UnsupportedDocumentError

SUPPORTED_SUFFIXES = {".docx", ".pdf", ".txt", ".md", ".markdown"}
_TEXT_ENCODINGS = ("utf-8-sig", "utf-16", "gb18030")


def read_text_file(path: str | Path) -> tuple[str, str]:
    source = Path(path)
    for encoding in _TEXT_ENCODINGS:
        try:
            return source.read_text(encoding=encoding), encoding
        except UnicodeDecodeError:
            continue
        except OSError as error:
            raise DocumentError("文本文件无法读取。") from error
    raise DocumentError("文本文件不是 UTF-8、UTF-16 或 GB18030 编码。")


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    try:
        with Path(path).open("rb") as handle:
            for block in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(block)
    except OSError as error:
        raise DocumentError("无法读取文件内容。") from error
    return digest.hexdigest()


def _table_rows(table: Table) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
        line = " | ".join(cells).rstrip()
        if line.strip(" |"):
            rows.append(line)
    return rows


def _body_blocks(document: Any) -> list[str]:
    blocks: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, document).text.strip()
            if text:
                blocks.append(text)
        elif child.tag.endswith("}tbl"):
            blocks.extend(_table_rows(Table(child, document)))
    return blocks


def _section_blocks(document: Any) -> tuple[list[str], list[str]]:
    headers: list[str] = []
    footers: list[str] = []
    for section in document.sections:
        for container, target in ((section.header, headers), (section.footer, footers)):
            for paragraph in container.paragraphs:
                text = paragraph.text.strip()
                if text:
                    target.append(text)
            for table in container.tables:
                target.extend(_table_rows(table))
    return _dedupe(headers), _dedupe(footers)


def _dedupe(values: list[str]) -> list[str]:
    result: list[str] = []
    for value in values:
        if value not in result:
            result.append(value)
    return result


def extract_docx(path: str | Path) -> str:
    try:
        document = Document(str(path))
    except Exception as error:
        raise DocumentError("选中的 Word 文档无法打开。") from error
    blocks = _body_blocks(document)
    headers, footers = _section_blocks(document)
    parts: list[str] = []
    if headers:
        parts.append("【页眉】\n" + "\n".join(headers))
    parts.extend(blocks)
    if footers:
        parts.append("【页脚】\n" + "\n".join(footers))
    text = "\n\n".join(part for part in parts if part.strip()).strip()
    if not text:
        raise DocumentError("Word 文档没有可提取的文本。")
    return text + "\n"


def extract_pdf_pages(path: str | Path) -> list[str]:
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise DocumentError("PDF 仍处于加密状态，请先解锁。")
        return [page.extract_text() or "" for page in reader.pages]
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError("选中的 PDF 无法读取。") from error


def is_scanned_pdf(pages: list[str]) -> bool:
    if not pages:
        return True
    visible = sum(len(re.sub(r"\s+", "", page)) for page in pages)
    return visible < max(40, len(pages) * 20)


def unlock_or_copy_pdf(
    source: str | Path,
    destination: str | Path,
    password: str | None,
) -> tuple[Path, Path | None, bool]:
    """Copy a PDF, unlocking it with a one-shot password when necessary."""
    source_path = Path(source)
    destination_path = Path(destination)
    try:
        reader = PdfReader(str(source_path))
    except Exception as error:
        raise DocumentError("选中的 PDF 无法打开。") from error
    if not reader.is_encrypted:
        try:
            shutil.copy2(source_path, destination_path)
        except OSError as error:
            raise DocumentError("PDF 无法复制到本地工作区。") from error
        return destination_path, None, False

    if not password:
        raise DocumentError("该 PDF 已加密；请使用“批量解锁 PDF”输入密码。")
    try:
        if not reader.decrypt(password):
            raise DocumentError("PDF 密码不正确。")
        encrypted_copy = destination_path.with_name(f"{destination_path.stem}.encrypted.pdf")
        shutil.copy2(source_path, encrypted_copy)
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        with destination_path.open("wb") as stream:
            writer.write(stream)
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError("PDF 无法在本地解锁。") from error
    return destination_path, encrypted_copy, True


def ensure_supported(path: str | Path) -> str:
    suffix = Path(path).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedDocumentError("只支持 Word、PDF、TXT 和 Markdown 文件。")
    return suffix
