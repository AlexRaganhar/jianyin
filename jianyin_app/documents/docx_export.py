"""Export redacted review text as a readable, text-reflowed DOCX."""

from __future__ import annotations

import re
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .markdown import PAGE_MARKER_RE

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)$")
_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")


def _set_run_font(run: object, name: str, size: int | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            style = document.styles[name]
            style.font.name = "微软雅黑"
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        except KeyError:
            continue


def _configure_page(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)


def _header_footer(document: Document, header_text: str) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(header_text)
    _set_run_font(run, "微软雅黑", 8)
    run.font.color.rgb = RGBColor(0x8A, 0x7A, 0x62)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = footer.add_run("第 ")
    _set_run_font(prefix, "宋体", 8)
    field_run = footer.add_run()
    _set_run_font(field_run, "宋体", 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.append(begin)
    field_run._r.append(instruction)
    field_run._r.append(separate)
    field_run._r.append(end)
    suffix = footer.add_run(" 页")
    _set_run_font(suffix, "宋体", 8)


def _add_page_caption(document: Document, page_number: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"第 {page_number} 页")
    _set_run_font(run, "宋体", 8)
    run.font.color.rgb = RGBColor(0x9A, 0x8C, 0x73)


def _join_lines(lines: list[str]) -> str:
    result = ""
    for line in lines:
        value = line.strip()
        if not value:
            continue
        if result and result[-1].isascii() and value[0].isascii():
            if result[-1].isalnum() or result[-1] in ".,;:!?":
                result += " "
        result += value
    return result


def _split_table_row(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|"):
        value = value[:-1]
    return [cell.strip() for cell in value.split("|")]


def _add_table(document: Document, rows: list[str]) -> None:
    parsed = []
    for _index, row in enumerate(rows):
        if _TABLE_SEPARATOR_RE.match(row):
            continue
        cells = _split_table_row(row)
        if cells:
            parsed.append(cells)
    if not parsed:
        return
    columns = max(len(row) for row in parsed)
    table = document.add_table(rows=len(parsed), cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(parsed):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.text = row[column_index] if column_index < len(row) else ""
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, "宋体", 10)
                    if row_index == 0:
                        run.bold = True


def _flush_paragraph(document: Document, buffer: list[str]) -> None:
    if not buffer:
        return
    text = _join_lines(buffer)
    if text:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        _set_run_font(run, "宋体", 11)
    buffer.clear()


def export_docx(
    markdown_text: str,
    output_path: str | Path,
    *,
    title: str = "脱敏文本",
    header_text: str = "笺隐 · 脱敏文本 · 请人工复核",
) -> Path:
    """Write a readable DOCX from redacted Markdown-like review text.

    This intentionally reflows text; it does not attempt to reproduce the source
    DOCX/PDF layout.
    """
    target = Path(output_path)
    document = Document()
    _configure_styles(document)
    _configure_page(document)
    _header_footer(document, header_text)
    document.core_properties.title = title
    document.core_properties.author = "笺隐"
    document.core_properties.comments = "文本重排导出；不含原始文件与映射。"

    lines = markdown_text.splitlines()
    buffer: list[str] = []
    table_buffer: list[str] = []
    current_page: int | None = None
    index = 0

    def flush_table() -> None:
        if table_buffer:
            _add_table(document, table_buffer)
            table_buffer.clear()

    while index < len(lines):
        line = lines[index]
        page_match = PAGE_MARKER_RE.match(line)
        if page_match:
            _flush_paragraph(document, buffer)
            flush_table()
            try:
                page_number = int(page_match.group(1))
            except (TypeError, ValueError):
                page_number = (current_page or 0) + 1
            if current_page is not None:
                document.add_page_break()
            current_page = page_number
            _add_page_caption(document, page_number)
            index += 1
            continue

        if not line.strip():
            _flush_paragraph(document, buffer)
            flush_table()
            index += 1
            continue

        if "|" in line:
            _flush_paragraph(document, buffer)
            table_buffer.append(line)
            index += 1
            continue
        flush_table()

        heading = _HEADING_RE.match(line)
        if heading:
            _flush_paragraph(document, buffer)
            level = min(len(heading.group(1)), 3)
            text = heading.group(2).strip()
            if text:
                paragraph = document.add_heading(level=level)
                run = paragraph.add_run(text)
                _set_run_font(run, "微软雅黑", 16 - level * 2)
            index += 1
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            _flush_paragraph(document, buffer)
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(bullet.group(1).strip())
            _set_run_font(run, "宋体", 11)
            index += 1
            continue

        numbered = _NUMBER_RE.match(line)
        if numbered:
            _flush_paragraph(document, buffer)
            paragraph = document.add_paragraph(style="List Number")
            run = paragraph.add_run(numbered.group(1).strip())
            _set_run_font(run, "宋体", 11)
            index += 1
            continue

        buffer.append(line)
        index += 1

    _flush_paragraph(document, buffer)
    flush_table()

    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f".{target.name}.{uuid4().hex}.tmp.docx")
    try:
        document.save(str(temporary))
        temporary.replace(target)
    except OSError as error:
        raise OSError(f"无法写入 DOCX：{target}") from error
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass
    return target
